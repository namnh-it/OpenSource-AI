from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[5]
TARGET = ROOT / "output" / "docx" / "So_tay_code_AI_Gateway_BTC.docx"


def set_font(run, name: str, size: float, *, bold: bool = False, color: str | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell_or_paragraph, fill: str) -> None:
    if hasattr(cell_or_paragraph, "_tc"):
        target = cell_or_paragraph._tc.get_or_add_tcPr()
    else:
        target = cell_or_paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    target.append(shd)


def set_cell_margin(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    set_font(run, "Aptos", 13)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_font(run, "Aptos", 13)


def add_code(doc: Document, label: str, code: str) -> None:
    label_p = doc.add_paragraph(style="Normal")
    label_p.paragraph_format.space_before = Pt(4)
    label_p.paragraph_format.space_after = Pt(2)
    label_run = label_p.add_run(label)
    set_font(label_run, "Aptos", 12, bold=True, color="1F4E79")

    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.18)
    p.paragraph_format.right_indent = Cm(0.18)
    p.paragraph_format.line_spacing = 1.0
    shade(p, "F3F6F8")
    run = p.add_run(code.rstrip())
    set_font(run, "Consolas", 9.5)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    if doc.tables:
        table.style = doc.tables[0].style
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margin(cell)
        shade(cell, "1F4E79")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        set_font(run, "Aptos", 11, bold=True, color="FFFFFF")
    no_split(table.rows[0])

    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        no_split(row)
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margin(cell)
            if row_index % 2 == 1:
                shade(cell, "EDF3F8")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(value)
            set_font(run, "Aptos", 11)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def append_agent_chapter(doc: Document) -> None:
    heading = doc.add_paragraph(style="Heading 1")
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(6)
    run = heading.add_run("15  Agent cơ bản với tool calling có giới hạn")
    set_font(run, "Aptos Display", 16, bold=True)

    add_body(
        doc,
        "Agent chỉ đáng dùng khi sản phẩm phải làm nhiều hơn một lần trả lời: nhận mục tiêu, chọn một hành động nhỏ, quan sát kết quả và mới chốt đầu ra. Nếu bài chỉ cần hỏi đáp, tóm tắt hoặc sinh một đoạn nội dung, gọi endpoint text thường sẽ rẻ, nhanh và ổn định hơn.",
    )
    add_body(
        doc,
        "Mã dưới đây bám theo khung đã có tại chung-khao/ai-starter. Thiết kế cố ý nhỏ: model chỉ được chọn tool trong allow-list; backend mới thực thi tool; số bước bị chặn cứng. Không biến agent thành quyền chạy shell, đọc file tùy ý hoặc gọi dịch vụ ngoài Gateway BTC.",
    )

    add_code(
        doc,
        "Luồng hoạt động",
        "goal + context\n"
        "      -> controller model chon tool\n"
        "      -> backend chi chay tool trong allow-list\n"
        "      -> ket qua duoc them vao messages + trace\n"
        "      -> final answer hoac dung khi dat max_steps",
    )

    sub = doc.add_paragraph(style="Heading 2")
    sub.add_run("15.1  Contract request và giới hạn bắt buộc")
    add_body(doc, "Request của agent dùng schema hẹp. Giá trị mặc định là 2 bước; tuyệt đối không để vòng lặp không giới hạn trong bài thi.")
    add_code(
        doc,
        "app/schemas.py",
        '''from typing import Literal\nfrom pydantic import BaseModel, Field\n\nAgentToolName = Literal["create_checklist", "draft_text", "generate_image"]\n\nclass AgentRunRequest(BaseModel):\n    goal: str = Field(min_length=3, max_length=4_000)\n    context: str | None = Field(default=None, max_length=12_000)\n    model: str | None = None\n    allowed_tools: list[AgentToolName] = Field(\n        default_factory=lambda: ["create_checklist", "draft_text"]\n    )\n    max_steps: int = Field(default=2, ge=1, le=4)\n    temperature: float = Field(default=0.2, ge=0, le=2)\n    max_completion_tokens: int = Field(default=1_200, ge=64, le=4_000)''',
    )
    add_table(
        doc,
        ["Trường", "Khuyến nghị khi demo", "Vì sao"],
        [
            ["goal", "Một mục tiêu hẹp, quan sát được", "Dễ chứng minh agent hoàn tất gì."],
            ["context", "Chỉ dữ kiện cần thiết; không có secret", "Giảm token và tránh lộ dữ liệu."],
            ["allowed_tools", "2 tool trước; image chỉ opt-in", "Giảm chi phí, độ trễ và lỗi."],
            ["max_steps", "2; tối đa 4 khi đã kiểm chứng", "Không treo UI hoặc cháy budget."],
            ["temperature", "0.2 cho workflow", "Output ổn định, dễ test lại."],
        ],
        [3.0, 6.3, 7.3],
    )

    sub = doc.add_paragraph(style="Heading 2")
    sub.add_run("15.2  Khai báo tool và allow-list")
    add_body(doc, "Tool definition là hợp đồng mà model nhìn thấy. Chỉ khai báo input tối thiểu và có schema chặt. Model không tự thực thi tool; nó chỉ trả tool call về backend.")
    add_code(
        doc,
        "app/agents/runner.py  -  tool definitions",
        '''AGENT_TOOL_DEFINITIONS = {\n    "create_checklist": {\n        "type": "function",\n        "function": {\n            "name": "create_checklist",\n            "description": "Turn concrete tasks into a numbered checklist.",\n            "parameters": {\n                "type": "object",\n                "properties": {\n                    "items": {\n                        "type": "array",\n                        "items": {"type": "string"},\n                        "maxItems": 12,\n                    }\n                },\n                "required": ["items"],\n                "additionalProperties": False,\n            },\n        },\n    },\n    "draft_text": {\n        "type": "function",\n        "function": {\n            "name": "draft_text",\n            "description": "Create one concise factual Vietnamese draft.",\n            "parameters": {\n                "type": "object",\n                "properties": {\n                    "instruction": {"type": "string"},\n                    "context": {"type": "string"},\n                },\n                "required": ["instruction"],\n                "additionalProperties": False,\n            },\n        },\n    },\n}''',
    )
    add_bullet(doc, "create_checklist chạy local, không phát sinh thêm một lần gọi model.")
    add_bullet(doc, "draft_text gọi lại Gateway BTC; dùng khi bản nháp thực sự làm đầu ra tốt hơn.")
    add_bullet(doc, "generate_image có sẵn trong ai-starter nhưng chỉ đưa vào allowed_tools khi ảnh là giá trị thấy được của sản phẩm.")

    sub = doc.add_paragraph(style="Heading 2")
    sub.add_run("15.3  Vòng chạy agent tối thiểu")
    add_body(doc, "Đây là logic chính trong AgentRunner. Mỗi vòng gửi messages + tools lên Gateway; nếu không còn tool call thì trả final_answer. Khi có tool call, kết quả được ghi vào trace và đưa lại vào messages trước vòng tiếp theo.")
    add_code(
        doc,
        "app/agents/runner.py  -  run loop rút gọn",
        '''async def run(self, request: AgentRunRequest) -> dict[str, Any]:\n    tools = [AGENT_TOOL_DEFINITIONS[name] for name in request.allowed_tools]\n    model = request.model or get_settings().ai_text_model\n    messages = [\n        {"role": "system", "content": (\n            "You are a bounded product-delivery agent. "\n            "Use a tool only when it moves the result forward. "\n            "Do not invent tool outputs or credentials. "\n            "Respond concisely in Vietnamese."\n        )},\n        {"role": "user", "content": f"Goal:\\n{request.goal}\\n\\nContext:\\n{request.context or '(none)'}"},\n    ]\n    trace = []\n\n    for step in range(1, request.max_steps + 1):\n        response = await self.gateway.chat({\n            "model": model, "messages": messages, "tools": tools,\n            "tool_choice": "auto", "temperature": request.temperature,\n            "max_completion_tokens": request.max_completion_tokens,\n        })\n        message = ((response.get("choices") or [{}])[0]).get("message") or {}\n        tool_calls = message.get("tool_calls") or []\n        if not tool_calls:\n            return {"status": "completed", "goal": request.goal,\n                    "final_answer": message.get("content") or "No final answer.",\n                    "model": response.get("model", model),\n                    "steps_used": step, "trace": trace}\n\n        messages.append({"role": "assistant", "content": message.get("content") or "",\n                         "tool_calls": tool_calls})\n        for call in tool_calls:\n            name = (call.get("function") or {}).get("name", "")\n            args = self._parse_arguments((call.get("function") or {}).get("arguments", "{}"))\n            result = await self._execute_tool(name, args, request)\n            trace.append({"step": step, "tool": name, "result": result})\n            messages.append({"role": "tool", "tool_call_id": call.get("id", f"step-{step}"),\n                             "content": json.dumps(result, ensure_ascii=False)})\n\n    return {"status": "max_steps_reached", "goal": request.goal,\n            "final_answer": "Da dung theo gioi han so buoc.",\n            "model": model, "steps_used": request.max_steps, "trace": trace}''',
    )

    sub = doc.add_paragraph(style="Heading 2")
    sub.add_run("15.4  Thực thi tool an toàn")
    add_body(doc, "Điểm quan trọng nhất là kiểm tra tên tool trước khi chạy. Nhờ đó model không thể yêu cầu một tool không được bật, kể cả khi nó sinh ra tool call lỗi hoặc bị prompt injection.")
    add_code(
        doc,
        "app/agents/runner.py  -  kiểm tra allow-list và tool local",
        '''async def _execute_tool(self, tool_name, arguments, request):\n    if tool_name not in request.allowed_tools:\n        return {"ok": False, "error": f"Tool '{tool_name}' is not allowed"}\n    if "_error" in arguments:\n        return {"ok": False, "error": arguments["_error"]}\n    if tool_name == "create_checklist":\n        return self._create_checklist(arguments)\n    if tool_name == "draft_text":\n        return await self._draft_text(arguments, request)\n    if tool_name == "generate_image":\n        return await self._generate_image(arguments)\n    return {"ok": False, "error": f"Unsupported tool '{tool_name}'"}\n\n@staticmethod\ndef _create_checklist(arguments):\n    items = arguments.get("items")\n    if not isinstance(items, list) or not items:\n        return {"ok": False, "error": "items must be a non-empty list"}\n    normalized = [str(item).strip()[:240] for item in items if str(item).strip()][:12]\n    return {"ok": True,\n            "checklist": [f"{index}. {item}" for index, item in enumerate(normalized, 1)]}''',
    )
    add_body(doc, "Không thêm shell, đọc/ghi file tùy ý, URL tùy ý, gửi email hay truy cập secret vào tool list. Một tool mới chỉ nên được thêm nếu người dùng nhìn thấy kết quả của nó trong UI và đội có thể giải thích rủi ro của nó.")

    sub = doc.add_paragraph(style="Heading 2")
    sub.add_run("15.5  Expose endpoint và gọi từ frontend")
    add_code(
        doc,
        "app/api/routes.py  -  FastAPI route",
        '''from app.agents.runner import AgentRunner\nfrom app.core.gateway import GatewayClient\nfrom app.schemas import AgentRunRequest\n\n@router.post("/agent/runs")\nasync def run_agent(request: AgentRunRequest) -> dict[str, Any]:\n    \"\"\"Run a short, auditable tool loop.\"\"\"\n    return await AgentRunner(GatewayClient()).run(request)''',
    )
    add_code(
        doc,
        "Frontend hoặc Postman  -  body mẫu",
        '''POST /api/ai/agent/runs\n{\n  "goal": "Tao checklist onboarding cho nguoi dung moi",\n  "context": "San pham la tro ly tai chinh ca nhan. Nguoi dung can 3 buoc dau tien.",\n  "model": "gpt-6-luna",\n  "max_steps": 2,\n  "allowed_tools": ["create_checklist", "draft_text"],\n  "temperature": 0.2,\n  "max_completion_tokens": 800\n}''',
    )
    add_body(doc, "UI nên hiển thị final_answer và trace dạng timeline: Bước 1 - tool nào đã dùng - kết quả tóm tắt; sau đó là câu trả lời cuối. Nếu model không gọi tool, vẫn hiển thị final_answer như một fallback hợp lệ; không retry vô hạn chỉ để ép agent dùng tool.")

    sub = doc.add_paragraph(style="Heading 2")
    sub.add_run("15.6  Test nhanh và script demo 20 giây")
    add_table(
        doc,
        ["Bước", "Thao tác", "Điều BGK nhìn thấy"],
        [
            ["1", "Nhập goal hẹp và context ngắn", "Agent có mục tiêu cụ thể, không nhận một lệnh mơ hồ."],
            ["2", "Bật 2 tool create_checklist và draft_text", "Tool được giới hạn theo đúng use case."],
            ["3", "Chạy tối đa 2 bước", "Có guardrail về tiền và thời gian phản hồi."],
            ["4", "Mở trace rồi final answer", "Chuỗi hành động quan sát được, có bằng chứng tool result."],
            ["5", "Nêu fallback và giới hạn", "Đội không tuyên bố agent tự trị hoặc quyết định thay con người."],
        ],
        [1.2, 7.1, 8.3],
    )
    add_body(doc, "Lời nói khi demo: “Agent của chúng em không chỉ trả lời một lần. Nó nhận một mục tiêu nhỏ, chỉ được chọn công cụ đã được giới hạn, quan sát kết quả rồi mới chốt output. Trace cho thấy từng bước, còn max_steps giúp kiểm soát thời gian và chi phí.”")
    add_bullet(doc, "Test mock đã có tại tests/test_agent_runner.py: kiểm tra trường hợp tool chạy rồi có final answer, và trường hợp tool nằm ngoài allow-list bị từ chối.")
    add_bullet(doc, "Chạy local trước demo: .\\.venv\\Scripts\\python.exe -m pytest tests/test_agent_runner.py -q")
    add_bullet(doc, "Nếu agent không tạo giá trị rõ hơn route text thường, tắt feature agent và demo luồng deterministic. Đó là quyết định tốt hơn việc gắn agent cho đủ tên gọi.")

    source_heading = doc.add_paragraph(style="Heading 2")
    source_heading.add_run("Mã nguồn tương ứng trong ai starter")
    add_bullet(doc, "app/agents/runner.py: tool definitions, AgentRunner, trace, giới hạn bước và implementation tool.")
    add_bullet(doc, "app/schemas.py: AgentRunRequest và AgentToolName.")
    add_bullet(doc, "app/api/routes.py: POST /api/ai/agent/runs.")
    add_bullet(doc, "docs/AGENT_PLAYBOOK.md: cách chọn use case, script thuyết trình và những điều không nên làm trong giờ thi.")


def main() -> None:
    if not TARGET.exists():
        raise FileNotFoundError(TARGET)
    doc = Document(TARGET)
    marker = "15  Agent cơ bản với tool calling có giới hạn"
    if any(paragraph.text.strip() == marker for paragraph in doc.paragraphs):
        raise RuntimeError("Agent chapter already exists; refusing to append it twice.")
    append_agent_chapter(doc)
    doc.save(TARGET)
    print(f"Updated {TARGET}")


if __name__ == "__main__":
    main()
