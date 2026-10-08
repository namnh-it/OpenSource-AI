from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "So_tay_Pipeline_LLM_Thi_Ngan.docx"
BLUE = "134A9B"
LIGHT_BLUE = "EAF2FF"
PALE_BLUE = "F5F9FF"
GRAY = "5B6573"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = tc_pr.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        tc_pr.append(shade)
    shade.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D7E0EC") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def style_run(run, *, size: float | None = None, bold: bool = False, color: str | None = None) -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(document, text: str = "", *, style: str | None = None, size: float | None = None, color: str | None = None, bold: bool = False, space_after: float = 6) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(space_after)
    if text:
        style_run(paragraph.add_run(text), size=size, color=color, bold=bold)


def add_bullets(document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        style_run(paragraph.add_run(item), size=10.5)


def add_table(document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, value in enumerate(headers):
        header_cells[index].width = Inches(widths[index])
        set_cell_shading(header_cells[index], BLUE)
        set_cell_border(header_cells[index], BLUE)
        paragraph = header_cells[index].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        style_run(run, size=9.2, bold=True, color="FFFFFF")
        header_cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        fill = "FFFFFF" if row_index % 2 == 0 else PALE_BLUE
        for index, value in enumerate(values):
            cells[index].width = Inches(widths[index])
            set_cell_shading(cells[index], fill)
            set_cell_border(cells[index])
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            style_run(run, size=9.2)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def heading(document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    style_run(paragraph.add_run(text), size=15 if level == 1 else 12, bold=True, color="000000")


def setup_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.color.rgb = RGBColor(0, 0, 0)

    if "Handbook Subtitle" not in document.styles:
        subtitle = document.styles.add_style("Handbook Subtitle", WD_STYLE_TYPE.PARAGRAPH)
        subtitle.font.name = "Aptos"
        subtitle.font.size = Pt(12)
        subtitle.font.color.rgb = RGBColor.from_string(GRAY)


def add_footer(document: Document) -> None:
    for section in document.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("AI Thực chiến - PiTA Team  |  Sổ tay Pipeline LLM")
        style_run(run, size=8.5, color=GRAY)


def main() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    setup_styles(document)
    add_footer(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(title.add_run("Sổ tay Pipeline LLM cho bài thi ngắn"), size=24, bold=True, color="000000")
    subtitle = document.add_paragraph(style="Handbook Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(subtitle.add_run("Chọn đúng lane, làm một happy path, giữ fallback để demo ổn định"), size=12, color=GRAY)
    document.add_paragraph()
    add_paragraph(
        document,
        "Dùng tài liệu này khi vừa nhận đề. Mục tiêu là chốt một hành động người dùng có kết quả nhìn thấy, "
        "không phải gom nhiều model vào cùng một demo.",
        size=11,
        space_after=10,
    )

    heading(document, "1. Chọn pipeline trong 60 giây")
    add_bullets(
        document,
        [
            "Viết một câu: ai làm gì, nhập gì, nhận gì và quyết định gì tiếp theo.",
            "Làm baseline không AI cho đúng hành động đó: rule, filter, fixture, công thức hoặc template.",
            "Chỉ chọn một lane AI tạo ra phần giá trị mà baseline chưa làm được.",
            "Chọn output nhỏ nhất có thể kiểm tra: JSON ngắn, một đoạn giải thích, một ảnh hoặc video 4 giây.",
            "Chốt fallback trước khi gọi model: fixture, kết quả baseline hoặc asset tĩnh.",
        ],
    )

    add_table(
        document,
        ["Người dùng cần", "Lane chính", "Không cần thêm"],
        [
            ["Giải thích, phân loại, tóm tắt, trích xuất, viết lại", "Text LLM", "Agent, RAG hay video chỉ để trang trí"],
            ["Trả lời dựa trên tài liệu thật", "Embedding + retrieval nhỏ", "RAG nếu chỉ có vài dòng form"],
            ["Một asset trực quan cho màn hình hoặc slide", "Image", "Sinh nhiều ảnh để chọn ngẫu nhiên"],
            ["Chuyển động là giá trị sản phẩm", "Video job bất đồng bộ", "Video nếu ảnh tĩnh đã đủ chứng minh luồng"],
            ["Chuỗi bước phụ thuộc nhau", "Agent có 1-2 tool", "Agent cho một câu trả lời đơn lẻ"],
        ],
        [2.55, 1.62, 2.75],
    )

    heading(document, "2. Pipeline chung")
    add_paragraph(
        document,
        "User action -> validate và normalize input -> baseline deterministic -> AI lane được chọn -> "
        "validate output -> render hoặc lưu artifact -> fallback khi lỗi, timeout hoặc vượt ngân sách -> audit.",
        size=10.5,
        bold=True,
        color=BLUE,
    )
    add_table(
        document,
        ["Ô cần chốt", "Câu hỏi cần trả lời"],
        [
            ["User action", "Một hành động hoàn chỉnh là gì?"],
            ["Input contract", "Field nào bắt buộc? Giới hạn ký tự, tệp và định dạng là gì?"],
            ["Baseline", "Không gọi AI thì vẫn xử lý hoặc hiển thị tối thiểu gì?"],
            ["AI lane", "Text, retrieval, image, video, audio hay agent?"],
            ["Output contract", "UI nhận chuỗi, JSON, artifact URL hay trạng thái job?"],
            ["Validation", "Kiểm tra schema, độ dài, policy hay confidence nào?"],
            ["Fallback", "Gateway lỗi thì user vẫn thấy gì?"],
        ],
        [1.72, 5.2],
    )

    heading(document, "3. Các lane có sẵn trong ai-starter")
    heading(document, "Text", level=2)
    add_paragraph(
        document,
        "Form hoặc fixture -> validate nghiệp vụ -> rule/baseline ngắn -> POST /api/ai/text -> "
        "kiểm tra output rỗng hoặc quá dài -> card kết quả, lý do và bước tiếp theo.",
    )
    add_bullets(
        document,
        [
            "Prompt gồm: vai trò, dữ kiện được phép dùng, định dạng output và điều không được bịa.",
            "Dùng nhiệt độ thấp cho phân loại và trích xuất. Nếu cần JSON, parse và validate ở backend trước khi đưa lên UI.",
            "Khởi đầu với gpt-6-luna; chỉ nâng model khi fixture đại diện không đạt.",
        ],
    )

    heading(document, "Retrieval", level=2)
    add_paragraph(
        document,
        "Docs đã chọn -> chunk cố định và metadata nguồn -> POST /api/ai/embeddings để index -> query embedding -> "
        "top 3 chunks -> POST /api/ai/text với chunks và yêu cầu nêu giới hạn.",
    )
    add_paragraph(
        document,
        "Chỉ dùng retrieval khi có corpus thật. Với vài dòng form hoặc ít hơn vài trang, truyền context trực tiếp thường nhanh và đáng tin hơn.",
        size=10.2,
        color=GRAY,
    )

    heading(document, "Image và video", level=2)
    add_table(
        document,
        ["Lane", "Luồng tối thiểu", "Fallback"],
        [
            ["Image", "Brief gồm subject, style, bố cục, tỉ lệ -> POST /api/ai/images, n = 1 -> preview artifact URL", "SVG/CSS hoặc ảnh tĩnh"],
            ["Video", "Storyboard 1 cảnh -> POST /api/ai/videos -> lưu video id -> polling -> download khi completed", "Poster tĩnh và trạng thái job rõ ràng"],
            ["Speech/STT", "Tệp hoặc text ngắn -> endpoint audio -> artifact hoặc transcript -> UI", "Text input/button thay thế"],
        ],
        [1.0, 3.8, 2.15],
    )
    add_paragraph(
        document,
        "Video luôn là job nền. Gửi một clip 4 giây sớm, làm việc khác trong lúc polling. Không để video chặn luồng UI và không chạy thêm clip ngẫu nhiên khi sắp demo.",
        size=10.2,
        color=GRAY,
    )

    heading(document, "Agent", level=2)
    add_paragraph(
        document,
        "Goal + context -> controller chọn tool trong allow-list -> tool result vào state -> controller trả lời cuối hoặc dừng theo giới hạn bước.",
    )
    add_bullets(
        document,
        [
            "Dùng POST /api/ai/agent/runs khi tool đầu tiên thực sự quyết định bước tiếp theo.",
            "Đặt max_steps = 2 trong demo. Tool mặc định: create_checklist và draft_text.",
            "Generate image phải được bật rõ ràng. Hiển thị trace dạng timeline; vẫn render final answer nếu model không gọi tool.",
        ],
    )

    heading(document, "4. Case card để luyện trước")
    add_paragraph(
        document,
        "Các case sau giúp luyện pipeline theo các hướng có tính khả thi và tác động xã hội từng được đội ghi chú. Chúng không phải đề chính thức.",
        size=10.2,
        color=GRAY,
    )
    add_table(
        document,
        ["Case", "Happy path", "AI lane", "Baseline và fallback"],
        [
            ["EdTech: lộ trình học", "Nhập mục tiêu, thời gian, sở thích; nhận kế hoạch 3 bước có lý do", "Text; agent checklist nếu thật sự cần", "Rule match theo mục tiêu/thời gian; fixture ba lộ trình"],
            ["FinTech: giải thích chi tiêu", "Nhập khoản chi; nhận nhóm chi tiêu và câu hỏi tự rà soát", "Rule tổng hợp + text", "Tổng theo nhóm bằng code; không tư vấn đầu tư hoặc tín dụng cá nhân"],
            ["Đô thị/logistics: triage phản ánh", "Nhập mô tả sự cố; nhận category, ưu tiên và phiếu tóm tắt", "Text classification/extraction", "Keyword mapping; nhãn cần người xác nhận"],
            ["Hành chính: chuẩn bị hồ sơ", "Chọn thủ tục; nhận checklist giấy tờ còn thiếu", "Text + retrieval nếu có quy định thật", "Checklist tĩnh; nêu nguồn và thời điểm dữ liệu"],
        ],
        [1.35, 2.3, 1.45, 1.85],
    )

    heading(document, "5. Nhịp triển khai 120 phút")
    add_table(
        document,
        ["Mốc", "Quyết định pipeline"],
        [
            ["0-16 phút", "Chốt user action, output contract, baseline, fallback và đúng một AI lane."],
            ["16-28 phút", "Smoke test một request thật; ghi model, request id và chi phí."],
            ["28-65 phút", "Hoàn thành happy path end-to-end; không thêm lane thứ hai."],
            ["65-83 phút", "Test normal, empty, invalid input và Gateway failure; hiển thị fallback."],
            ["83-108 phút", "Chỉ thêm agent hoặc media nếu xuất hiện trong demo và có bằng chứng giá trị."],
            ["108-120 phút", "Dừng gọi model mới trừ blocker; đóng gói demo, fixture và câu chuyện trình bày."],
        ],
        [1.25, 5.7],
    )

    heading(document, "6. Mẫu điền nhanh trước khi code")
    add_table(
        document,
        ["Trường", "Nội dung cần chốt"],
        [
            ["User action", ""],
            ["Input contract", ""],
            ["Kết quả user nhìn thấy", ""],
            ["Quyết định tiếp theo của user", ""],
            ["Baseline không AI", ""],
            ["AI lane duy nhất", ""],
            ["Model và max output/call", ""],
            ["Output validation", ""],
            ["Fallback hoặc fixture", ""],
            ["Một câu demo", ""],
        ],
        [2.25, 4.7],
    )
    add_paragraph(
        document,
        "Nếu chưa điền được kết quả user nhìn thấy và fallback, chưa gọi model. Nếu agent, image hoặc video không thay đổi câu demo, bỏ chúng.",
        size=10.5,
        bold=True,
        color=BLUE,
        space_after=0,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
