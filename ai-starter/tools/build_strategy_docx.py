from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
OUTPUT = ROOT / "Chien_luoc_thi_AITC_PiTA_Team.docx"

NAVY = "1F4E78"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9E2F3"
GRID = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_borders(cell, color: str = GRID) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 110, start: int = 110, bottom: int = 110, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_font(run, *, size: float | None = None, bold: bool | None = None, color: RGBColor = BLACK) -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.color.rgb = color
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_text(paragraph, text: str, *, bold: bool = False, size: float | None = None) -> None:
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold)


def style_paragraph(paragraph, *, after: float = 6, before: float = 0, line: float = 1.12) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(after)
    fmt.space_before = Pt(before)
    fmt.line_spacing = line


def add_body(doc: Document, text: str, *, after: float = 6) -> None:
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, after=after)
    add_text(paragraph, text, size=10.2)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        style_paragraph(paragraph, after=3)
        add_text(paragraph, item, size=10.0)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(13 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_font(run, size=15 if level == 1 else 11.5, bold=True)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    set_repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(paragraph, after=0, line=1.0)
        run = paragraph.add_run(header)
        set_font(run, size=8.8, bold=True, color=RGBColor(255, 255, 255))

    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            style_paragraph(paragraph, after=0, line=1.04)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            set_font(run, size=8.7)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    set_font(run, size=8.5, color=RGBColor(89, 89, 89))
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = BLACK

    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.color.rgb = BLACK

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(footer, after=0)
    add_text(footer, "PiTA Team  |  Chiến lược thi AITC  |  Trang ", size=8.5)
    add_page_field(footer)


def build_document() -> None:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(7)
    add_text(title, "Kế hoạch chiến lược thi AITC PiTA Team", bold=True, size=24)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.paragraph_format.space_after = Pt(16)
    add_text(subtitle, "Định hướng idea và PoC sản phẩm demo được trong thời gian giới hạn", size=12.5)

    intro = doc.add_paragraph()
    style_paragraph(intro, after=8, line=1.18)
    add_text(intro, "Kết luận điều hành. ", bold=True, size=10.4)
    add_text(
        intro,
        "Cơ hội tốt nhất để tạo khác biệt không nằm ở số lượng tính năng hay model đắt tiền. "
        "Đội cần chọn một vấn đề đủ cụ thể, chứng minh được giá trị bằng một luồng người dùng hoàn chỉnh, "
        "và trình diễn PoC ổn định ngay cả khi dịch vụ AI chậm hoặc lỗi.",
        size=10.4,
    )
    add_body(
        doc,
        "Tài liệu này là playbook dùng trước và trong cuộc thi. Khi nhận đề, đội thay các giả định bằng yêu cầu và tiêu chí chấm chính thức; không coi các ví dụ trong tài liệu là yêu cầu của Ban tổ chức.",
        after=10,
    )

    add_heading(doc, "1  Mục tiêu thắng cuộc và nguyên tắc lựa chọn", 1)
    add_body(
        doc,
        "Mục tiêu của PiTA Team là biến một idea có ý nghĩa thành bằng chứng dễ kiểm tra: người dùng đưa vào dữ liệu, hệ thống tạo một kết quả hữu ích, và đội có thể giải thích phần nào là AI, phần nào là baseline, cùng giới hạn của giải pháp.",
    )
    add_table(
        doc,
        ["Tiêu chí", "Câu hỏi quyết định", "Dấu hiệu chọn"],
        [
            ["Vấn đề", "Ai ra quyết định gì sau khi dùng sản phẩm?", "Có một người dùng và hành động rõ ràng."],
            ["Dữ liệu", "Input nào có thể có ngay khi thi?", "Có fixture hoặc dữ liệu mẫu hợp lệ."],
            ["AI", "Baseline không làm được phần nào?", "AI tạo giá trị thấy được, không chỉ gắn nhãn AI."],
            ["Demo", "Có kể trong 90 giây không?", "Một luồng chính, kết quả trực quan, có fallback."],
            ["Khả thi", "Ba người có thể chạy ổn định trong 2 giờ?", "Không phụ thuộc đào tạo dài, tích hợp phức tạp."],
        ],
        [2.3, 6.3, 7.0],
    )
    add_body(doc, "Quy tắc cắt scope: nếu một tính năng không giúp người dùng hoàn thành hành động chính hoặc không giúp hội đồng hiểu giá trị trong demo, để ngoài phạm vi.")

    add_heading(doc, "2  Cấu trúc một PoC có sức thuyết phục", 1)
    add_body(doc, "PoC cần là sản phẩm nhỏ có thể kiểm chứng, không phải bản mô tả giải pháp lớn. Mỗi đề chỉ chọn một lát cắt người dùng từ đầu đến cuối.")
    add_table(
        doc,
        ["Thành phần", "Mức tối thiểu phải có", "Minh chứng khi demo"],
        [
            ["Input", "Form, file hoặc dữ liệu mẫu có validation.", "Người xem thấy dữ liệu được đưa vào."],
            ["Baseline", "Rule, fixture hoặc phép tính không AI.", "Có điểm so sánh và phương án khi AI lỗi."],
            ["Năng lực AI", "Đúng một tác vụ: trích xuất, phân loại, sinh nội dung, thị giác hoặc retrieval.", "Kết quả có ích và giải thích được."],
            ["Output", "Kết quả có thể hành động tiếp, không chỉ văn bản chung chung.", "Có trạng thái thành công, rỗng và lỗi."],
            ["Bằng chứng", "Test hoặc dữ liệu mẫu đã chạy.", "Nêu rõ điều đã kiểm tra và giới hạn."],
        ],
        [2.3, 6.4, 6.9],
    )
    add_bullets(
        doc,
        [
            "Ưu tiên giải pháp 'narrow but deep': một quyết định được hỗ trợ tốt hơn nhiều màn hình nửa hoàn chỉnh.",
            "Giữ secret ở backend; frontend chỉ gọi API local của sản phẩm.",
            "Thiết kế trạng thái loading, timeout và fallback trước khi làm hiệu ứng hoặc asset phụ.",
        ],
    )

    add_heading(doc, "3  Vận hành đội trong sprint 120 phút", 1)
    add_body(doc, "Ba vai trò làm song song nhưng cùng chịu trách nhiệm cho một happy path. Mọi mốc dưới đây là checkpoint cắt scope, không phải danh sách tính năng bắt buộc.")
    add_table(
        doc,
        ["Thời gian", "Mục tiêu chung", "AI Engineer", "Data Scientist", "Software Engineer"],
        [
            ["0 - 16", "Chốt brief và một lát cắt", "Khả năng model, secret, smoke plan", "Input, baseline, metric", "Wireframe và starter nhỏ"],
            ["16 - 45", "Happy path chạy local", "Một integration AI", "Fixture và kiểm tra output", "UI, API, validation"],
            ["45 - 80", "Đủ tin cậy để demo", "Fallback và audit cost", "Normal, empty, adverse test", "Trạng thái lỗi, responsive"],
            ["80 - 108", "Câu chuyện và sửa blocker", "Triage lỗi có log", "Evidence và limitation", "Demo path, README"],
            ["108 - 120", "Gói nộp", "Không gọi model mới trừ blocker", "Xác nhận dữ liệu minh họa", "Run từ terminal mới"],
        ],
        [1.6, 3.4, 3.8, 3.8, 3.8],
    )
    add_body(doc, "Phân bổ năng lực tham chiếu: 45 phút web và product, 35 phút AI và backend, 20 phút QA và fallback, 15 phút demo và tài liệu, 5 phút buffer. Nếu không có đủ ba người, giữ thứ tự ưu tiên: happy path, fallback, test, rồi mới polish.")

    add_heading(doc, "4  Chiến lược model và ngân sách AI", 1)
    add_body(doc, "Ngân sách 50 USD nên được xem là trần cần bảo vệ, không phải mục tiêu phải dùng hết. Ghi chi phí response của Gateway và tách chi phí authoring hỗ trợ code khỏi chi phí runtime của sản phẩm.")
    add_table(
        doc,
        ["Ngăn chi", "Trần", "Mục đích", "Quy tắc"],
        [
            ["Smoke test", "$4", "Xác nhận 1 capability hoặc model.", "Một request đại diện, không thử ngẫu nhiên."],
            ["Hỗ trợ code", "$6", "Review diff hoặc triage bug có log.", "Không gửi toàn repo hoặc secret."],
            ["Runtime lõi", "$25", "Luồng người dùng chính đã chốt.", "Dừng khi vượt 70% ngăn chi."],
            ["Asset demo", "$7", "Ảnh, video, audio thực sự xuất hiện khi demo.", "Sinh ít, có storyboard trước."],
            ["Dự phòng", "$8", "Blocker tái lập và validation cuối.", "Chỉ người điều phối cho phép dùng."],
        ],
        [2.5, 1.6, 6.3, 4.0],
    )
    add_table(
        doc,
        ["Tác vụ", "Mặc định", "Chỉ nâng cấp khi"],
        [
            ["Review code và triage", "gpt-6-luna, effort low", "Lỗi vẫn tái lập sau patch nhỏ: 1 call gpt-6.1-sol."],
            ["Text đơn giản", "gpt-6-luna, effort none", "Fixture đại diện cho thấy chưa đạt chất lượng."],
            ["Copy hoặc tóm tắt", "deepseek-flash, tắt thinking", "Cần reasoning thật sự."],
            ["Ảnh hoặc video", "SVG/CSS local trước; Veo Lite 4 giây khi cần", "Asset là phần thiết yếu của demo."],
        ],
        [3.0, 4.9, 6.5],
    )
    add_body(doc, "Không dùng model cao nhất chỉ để tạo cảm giác 'mạnh'. Với task code nhỏ, context và test log tốt thường đem lại lợi ích lớn hơn đổi sang model đắt tiền.")

    add_heading(doc, "5  Kịch bản demo và câu chuyện idea", 1)
    add_body(doc, "Demo hiệu quả trả lời bốn câu: ai đang gặp vấn đề, họ làm gì trong sản phẩm, AI thay đổi quyết định như thế nào, và đội đã kiểm tra điều gì.")
    add_table(
        doc,
        ["Thời lượng", "Nội dung", "Bằng chứng phải hiện trên màn hình"],
        [
            ["0 - 15 giây", "Bối cảnh, người dùng, quyết định cần hỗ trợ.", "Một tình huống cụ thể, không mở đầu bằng kiến trúc."],
            ["15 - 55 giây", "Người dùng nhập/chọn dữ liệu và hoàn thành một hành động.", "Input, xử lý, output có thể hành động."],
            ["55 - 75 giây", "Phân biệt baseline và phần AI tạo giá trị.", "Lý do kết quả, nguồn hoặc mức tin cậy nếu có."],
            ["75 - 90 giây", "Giới hạn, fallback, bước tiếp theo.", "Ứng xử khi input rỗng hoặc Gateway lỗi."],
        ],
        [2.2, 6.0, 6.2],
    )
    add_bullets(
        doc,
        [
            "Không khẳng định tác động xã hội, độ chính xác hoặc tiết kiệm thời gian nếu chưa có dữ liệu hoặc phép đo.",
            "Nếu dữ liệu demo là mẫu, nói rõ đó là dữ liệu mẫu và mô tả cách kiểm tra tiếp theo.",
            "Một screenshot hoặc output thật có giá trị hơn một sơ đồ kiến trúc dày đặc.",
        ],
    )

    add_heading(doc, "6  Checklist trước khi trình bày hoặc nộp", 1)
    add_table(
        doc,
        ["Nhóm", "Kiểm tra bắt buộc"],
        [
            ["Chạy được", "Có một lệnh run, health check và đường demo đã chạy lại từ terminal mới."],
            ["AI", "Model, input limit, fallback, request cost và lỗi Gateway đã được kiểm tra."],
            ["Dữ liệu", "Input rỗng, input lỗi và một adverse case có kết quả rõ ràng."],
            ["Bảo mật", "Không có API key trong frontend, source, prompt, screenshot hoặc log."],
            ["Bằng chứng", "Lưu test đã chạy, fixture, screenshot và giới hạn đã biết."],
            ["Trình bày", "Demo dưới 90 giây, chữ dễ đọc, chỉ nêu claim có bằng chứng."],
        ],
        [3.0, 13.0],
    )
    add_heading(doc, "7  Hành động chuẩn bị trước ngày thi", 1)
    add_bullets(
        doc,
        [
            "Mỗi thành viên clone và chạy thành công ai-starter bằng máy của mình; kiểm tra Python, môi trường và API health trước giờ thi.",
            "Tạo sẵn thư mục task theo cấu trúc brief, solution, technical, experiment và pitch để ghi quyết định thay vì nhớ trong chat.",
            "Chuẩn bị 2 - 3 fixture trung tính cho text, ảnh hoặc file để test UI khi chưa có dữ liệu đề.",
            "Luyện một lần flow: nhận đề, chọn lát cắt, smoke test capability, hoàn thiện happy path, demo và review cuối.",
        ],
    )
    add_body(
        doc,
        "Tham chiếu kỹ thuật: AI Thực Chiến Gateway, bảng giá model và hướng dẫn model OpenAI/DeepSeek, kiểm tra ngày 05 tháng 10 năm 2026. Danh sách model, giá và giới hạn có thể thay đổi; xác nhận lại trước mỗi lần triển khai.",
        after=0,
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
